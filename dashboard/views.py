from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.http import HttpResponseForbidden, JsonResponse, HttpResponse
from django.utils import timezone
from django.db.models import Sum, Count, Avg
from datetime import timedelta
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from core.decorators import solo_gerencia
from pedidos.models import Pedido
from presupuestos.models import Presupuesto
from productos.models import ProductoCotizado
from clientes.models import Cliente
from .models import ObjetivoVentas
from .metrics import panel_horas


@login_required
@solo_gerencia
def dashboard(request):

    hoy = timezone.now().date()

    # ── Período ──────────────────────────────────────────────────────────────
    periodo = request.GET.get('periodo', 'mes')
    if periodo == 'semana':
        fecha_desde = hoy - timedelta(days=7)
    elif periodo == 'trimestre':
        fecha_desde = hoy - timedelta(days=90)
    else:
        fecha_desde = hoy - timedelta(days=30)

    # ── Pedidos en curso ─────────────────────────────────────────────────────
    estados_en_curso = ['No iniciado', 'En proceso']
    pedidos_en_curso = Pedido.objects.filter(
        estado__in=estados_en_curso,
        bloqueado_cancelado=False
    )
    total_saldo_pendiente = pedidos_en_curso.aggregate(
        total=Sum('saldo'))['total'] or 0

    # ── Pedidos completados en el período ────────────────────────────────────
    estados_completados = ['Terminado (falta pago)', 'Terminado y pagado']
    pedidos_completados = Pedido.objects.filter(
        estado__in=estados_completados,
        created__date__gte=fecha_desde  # usamos created ya que updated es auto_now_add
    )
    total_cobrado = 0
    for p in pedidos_completados:
        if not p.saldo:
            total_cobrado += p.precio
        elif p.senia:
            total_cobrado += p.senia

    ticket_promedio = (
        pedidos_completados.aggregate(avg=Avg('precio'))['avg'] or 0
    )

    # ── Pedidos por estado ───────────────────────────────────────────────────
    pedidos_por_estado = (
        Pedido.objects.filter(bloqueado_cancelado=False)
        .values('estado')
        .annotate(cantidad=Count('numero'))
        .order_by('estado')
    )

    # ── Presupuestos últimos 15 días ─────────────────────────────────────────
    hace_15_dias = hoy - timedelta(days=15)
    presupuestos_recientes = Presupuesto.objects.filter(
        created__date__gte=hace_15_dias)
    total_presupuestos = presupuestos_recientes.count()
    numeros_presupuesto = presupuestos_recientes.values_list(
        'numero', flat=True)
    presupuestos_convertidos = Pedido.objects.filter(
        presupuesto__in=numeros_presupuesto).count()
    tasa_conversion = (
        round((presupuestos_convertidos / total_presupuestos) * 100)
        if total_presupuestos > 0 else 0
    )

    # ── Productos más vendidos ───────────────────────────────────────────────
    presupuestos_con_pedido = Pedido.objects.filter(
        presupuesto__isnull=False,
        created__date__gte=fecha_desde
    ).values_list('presupuesto', flat=True)

    productos_vendidos = (
        ProductoCotizado.objects
        .filter(presupuesto__numero__in=presupuestos_con_pedido)
        .values('insumo__nombre')
        .annotate(
            veces_vendido=Count('id'),
            total_unidades=Sum('cantidad'),
            total_facturado=Sum('resultado')
        )
        .order_by('-veces_vendido')[:5]
    )

    # ── Top 5 clientes (6 meses) ─────────────────────────────────────────────
    hace_6_meses = hoy - timedelta(days=180)
    top_clientes = (
        Pedido.objects
        .filter(cliente__isnull=False, bloqueado_cancelado=False,
                created__date__gte=hace_6_meses)
        .values('cliente__id', 'cliente__nombre', 'cliente__negocio')
        .annotate(cantidad_pedidos=Count('numero'), total_comprado=Sum('precio'))
        .order_by('-total_comprado')[:5]
    )

    # ── Próximos pedidos con fecha de entrega ────────────────────────────────
    proximos_pedidos = (
        Pedido.objects
        .filter(fecha_entrega__gte=hoy, bloqueado_cancelado=False,
                estado__in=estados_en_curso)
        .order_by('fecha_entrega')[:8]
    )

    # ── Objetivos de ventas ──────────────────────────────────────────────────
    todos_objetivos = ObjetivoVentas.objects.all()
    objetivos_con_progreso = []

    for obj in todos_objetivos:
        pedidos_obj = Pedido.objects.filter(
            estado__in=estados_completados,
            created__date__gte=obj.fecha_desde,
            created__date__lte=obj.fecha_hasta
        )
        cobrado = 0
        for p in pedidos_obj:
            if not p.saldo:
                cobrado += p.precio
            elif p.senia:
                cobrado += p.senia

        progreso = min(round((cobrado / float(obj.monto_objetivo)) * 100), 100) \
            if obj.monto_objetivo > 0 else 0
        activo = obj.fecha_desde <= hoy <= obj.fecha_hasta

        objetivos_con_progreso.append({
            'id': obj.id,
            'nombre': obj.nombre,
            'monto_objetivo': obj.monto_objetivo,
            'fecha_desde': obj.fecha_desde,
            'fecha_hasta': obj.fecha_hasta,
            'cobrado': cobrado,
            'progreso': progreso,
            'activo': activo,
        })

    objetivos_vigentes = [o for o in objetivos_con_progreso if o['activo']]

    # ── Panel de horas / punto de equilibrio (Fase 5) ────────────────────────
    # Mes en curso, independiente del filtro de período de arriba.
    horas = panel_horas(hoy)

    context = {
        'autorizado': request.session.get('autorizado'),
        'usuario': request.session.get('usuario_nombre'),
        'img': request.session.get('img'),
        'periodo': periodo,
        'horas': horas,
        'pedidos_en_curso_count': pedidos_en_curso.count(),
        'total_saldo_pendiente': total_saldo_pendiente,
        'pedidos_completados_count': pedidos_completados.count(),
        'total_cobrado': total_cobrado,
        'ticket_promedio': ticket_promedio,
        'pedidos_por_estado': list(pedidos_por_estado),
        'total_presupuestos': total_presupuestos,
        'presupuestos_convertidos': presupuestos_convertidos,
        'tasa_conversion': tasa_conversion,
        'productos_vendidos': list(productos_vendidos),
        'top_clientes': list(top_clientes),
        'proximos_pedidos': proximos_pedidos,
        'hoy': hoy,
        'objetivos_vigentes': objetivos_vigentes,
        'todos_objetivos': objetivos_con_progreso,
        'hay_objetivos': len(objetivos_con_progreso) > 0,
    }
    return render(request, 'dashboard/dashboard.html', context)


@solo_gerencia
@login_required
def guardar_objetivo(request):

    if request.method == 'POST':
        nombre = request.POST.get('nombre')
        monto = request.POST.get('monto_objetivo')
        fecha_desde = request.POST.get('fecha_desde')
        fecha_hasta = request.POST.get('fecha_hasta')
        obj_id = request.POST.get('objetivo_id')

        if obj_id:
            obj = get_object_or_404(ObjetivoVentas, pk=obj_id)
            obj.nombre = nombre
            obj.monto_objetivo = monto
            obj.fecha_desde = fecha_desde
            obj.fecha_hasta = fecha_hasta
            obj.save()
        else:
            ObjetivoVentas.objects.create(
                nombre=nombre,
                monto_objetivo=monto,
                fecha_desde=fecha_desde,
                fecha_hasta=fecha_hasta
            )
    return redirect('/dashboard/')


@solo_gerencia
@login_required
def eliminar_objetivo(request, pk):

    obj = get_object_or_404(ObjetivoVentas, pk=pk)
    obj.delete()
    return redirect('/dashboard/')


@solo_gerencia
@login_required
def exportar_excel(request, tipo):

    hoy = timezone.now().date()
    periodo = request.GET.get('periodo', 'mes')
    if periodo == 'semana':
        fecha_desde = hoy - timedelta(days=7)
    elif periodo == 'trimestre':
        fecha_desde = hoy - timedelta(days=90)
    else:
        fecha_desde = hoy - timedelta(days=30)

    wb = openpyxl.Workbook()
    ws = wb.active

    header_fill = PatternFill("solid", fgColor="1a1a2e")
    header_font = Font(color="FFFFFF", bold=True)
    center = Alignment(horizontal="center")

    if tipo == 'pedidos':
        ws.title = "Pedidos"
        headers = ['#', 'Cliente', 'Descripción', 'Estado',
                   'Total', 'Seña', 'Saldo', 'Encargado', 'Fecha entrega']
        qs = Pedido.objects.filter(
            created__date__gte=fecha_desde
        ).select_related('cliente', 'encargado').order_by('-numero')

        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center

        for row, p in enumerate(qs, 2):
            ws.cell(row=row, column=1, value=p.numero)
            ws.cell(row=row, column=2, value=str(
                p.cliente) if p.cliente else '')
            ws.cell(row=row, column=3, value=p.descripcion)
            ws.cell(row=row, column=4, value=p.estado)
            ws.cell(row=row, column=5, value=p.precio)
            ws.cell(row=row, column=6, value=p.senia or 0)
            ws.cell(row=row, column=7, value=p.saldo or 0)
            ws.cell(row=row, column=8, value=p.nombre_encargado)
            ws.cell(row=row, column=9, value=str(
                p.fecha_entrega) if p.fecha_entrega else '')

    elif tipo == 'presupuestos':
        ws.title = "Presupuestos"
        headers = ['#', 'Cliente', 'Total',
                   'Descuento', 'Seña', 'Saldo', 'Fecha']
        qs = Presupuesto.objects.filter(
            created__date__gte=fecha_desde
        ).select_related('cliente').order_by('-numero')

        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center

        for row, p in enumerate(qs, 2):
            ws.cell(row=row, column=1, value=p.numero)
            ws.cell(row=row, column=2, value=str(
                p.cliente) if p.cliente else '')
            ws.cell(row=row, column=3, value=float(p.total))
            ws.cell(row=row, column=4, value=float(p.desc_plata))
            ws.cell(row=row, column=5, value=p.seña or 0)
            ws.cell(row=row, column=6, value=p.saldo or 0)
            ws.cell(row=row, column=7, value=str(p.created.date()))

    # Autoajustar columnas
    for col in ws.columns:
        max_len = max((len(str(c.value or '')) for c in col), default=10)
        ws.column_dimensions[get_column_letter(
            col[0].column)].width = min(max_len + 4, 40)

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{tipo}_{periodo}_{hoy}.xlsx"'
    wb.save(response)
    return response
